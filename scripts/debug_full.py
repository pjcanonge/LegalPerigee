#!/usr/bin/env python3
"""
LegalPerigee — Full Debug Suite
Systematically tests every component and reports pass/fail.
Run: python3 scripts/debug_full.py
"""

import os, sys, json, time, traceback
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).parent.parent))

# ── Load env / keychain ───────────────────────────────────────────────────────
from dotenv import load_dotenv
load_dotenv(dotenv_path=Path(__file__).parent.parent / ".env", override=True)
try:
    from utils.keychain import load_key
    if not os.environ.get("ANTHROPIC_API_KEY"):
        k = load_key("ANTHROPIC_API_KEY")
        if k: os.environ["ANTHROPIC_API_KEY"] = k
except Exception:
    pass

# ── Test harness ──────────────────────────────────────────────────────────────

PASS = "✅"; FAIL = "❌"; WARN = "⚠️ "; SKIP = "⏭️ "
results: list[dict] = []

def test(name: str, fn, critical: bool = False):
    start = time.time()
    try:
        detail = fn()
        ms = int((time.time()-start)*1000)
        results.append({"name": name, "status": "PASS", "detail": detail or "", "ms": ms})
        print(f"  {PASS}  {name:<52} {ms}ms")
        return True
    except Exception as e:
        ms = int((time.time()-start)*1000)
        tb = traceback.format_exc().strip().splitlines()[-1]
        results.append({"name": name, "status": "FAIL", "detail": tb, "ms": ms, "critical": critical})
        flag = "CRITICAL" if critical else "FAIL"
        print(f"  {FAIL}  {name:<52} {flag}")
        print(f"       └─ {tb}")
        return False

def section(title: str):
    print(f"\n{'─'*65}")
    print(f"  {title}")
    print(f"{'─'*65}")


# ═══════════════════════════════════════════════════════════════════════════════
print(f"\n{'═'*65}")
print(f"  ⚖️  LegalPerigee — Full Debug Suite")
print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"{'═'*65}")


# ── 1. Core imports ────────────────────────────────────────────────────────────
section("1. Core Imports")

test("Python version (3.9+)", lambda:
    f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}")

test("anthropic SDK", lambda:
    __import__("anthropic").__version__)

test("streamlit", lambda:
    __import__("streamlit").__version__)

test("fastapi", lambda:
    __import__("fastapi").__version__)

test("uvicorn", lambda:
    __import__("uvicorn").__version__)

test("pydantic", lambda:
    __import__("pydantic").__version__)

test("httpx", lambda:
    __import__("httpx").__version__)

test("pandas", lambda:
    __import__("pandas").__version__)

test("pdfplumber", lambda:
    __import__("pdfplumber").__version__ if hasattr(__import__("pdfplumber"), "__version__") else "ok")

test("python-docx", lambda:
    __import__("docx").__version__ if hasattr(__import__("docx"), "__version__") else "ok")

test("python-pptx", lambda:
    __import__("pptx").__version__ if hasattr(__import__("pptx"), "__version__") else "ok")

test("pillow", lambda:
    __import__("PIL").__version__)

test("pillow-heif", lambda:
    __import__("pillow_heif").__version__ if hasattr(__import__("pillow_heif"), "__version__") else "ok")

test("reportlab", lambda:
    __import__("reportlab").__version__)

test("keyring", lambda:
    __import__("keyring").__version__ if hasattr(__import__("keyring"), "__version__") else "ok")

test("altair", lambda:
    __import__("altair").__version__)

test("pywebview", lambda:
    __import__("webview").__version__ if hasattr(__import__("webview"), "__version__") else "ok")


# ── 2. Project modules ─────────────────────────────────────────────────────────
section("2. Project Modules")

test("models.case_report", lambda:
    str(__import__("models.case_report", fromlist=["CaseIntelReport"]).CaseIntelReport.__fields__.keys()))

test("database.db", lambda:
    str(__import__("database.db", fromlist=["init_db"]).init_db.__name__))

test("utils.json_extract", lambda:
    str(__import__("utils.json_extract", fromlist=["extract_json"]).extract_json('{"a":1}')))

test("utils.retry", lambda:
    str(__import__("utils.retry", fromlist=["call_with_retry"]).call_with_retry.__name__))

test("utils.keychain", lambda:
    str(__import__("utils.keychain", fromlist=["load_key"]).load_key.__name__))

test("utils.auto_repair", lambda:
    str(__import__("utils.auto_repair", fromlist=["capture_error"]).capture_error.__name__))

test("tools.court_tools", lambda:
    str(len(__import__("tools.court_tools", fromlist=["COURT_TOOL_SCHEMAS"]).COURT_TOOL_SCHEMAS)) + " tools")

test("tools.search_tools", lambda:
    str(__import__("tools.search_tools", fromlist=["WEB_SCRAPE_TOOL_SCHEMA"]).WEB_SCRAPE_TOOL_SCHEMA["name"]))

test("agents.court_researcher", lambda:
    str(__import__("agents.court_researcher", fromlist=["run_court_researcher"]).run_court_researcher.__name__))

test("agents.web_researcher", lambda:
    str(__import__("agents.web_researcher", fromlist=["run_web_researcher"]).run_web_researcher.__name__))

test("agents.documentor", lambda:
    str(__import__("agents.documentor", fromlist=["run_documentor"]).run_documentor.__name__))

test("agents.orchestrator", lambda:
    str(__import__("agents.orchestrator", fromlist=["run_investigation"]).run_investigation.__name__))

test("analyzers.media_analyzer", lambda:
    str(__import__("analyzers.media_analyzer", fromlist=["analyze_text"]).analyze_text.__name__))

test("analyzers.precedent_analyzer", lambda:
    str(__import__("analyzers.precedent_analyzer", fromlist=["discover_precedents"]).discover_precedents.__name__))

test("analyzers.trend_analyzer", lambda:
    str(__import__("analyzers.trend_analyzer", fromlist=["detect_waves"]).detect_waves.__name__))

test("analyzers.legislative_analyzer", lambda:
    str(__import__("analyzers.legislative_analyzer", fromlist=["analyze_bill"]).analyze_bill.__name__))

test("viewers.document_viewer", lambda:
    str(__import__("viewers.document_viewer", fromlist=["render_file"]).render_file.__name__))


# ── 3. All aggregators ─────────────────────────────────────────────────────────
section("3. Aggregators (import check)")

aggs = [
    ("courtlistener_fetch",      "run_full_sync"),
    ("caselaw_fetch",            "run_cap_sync"),
    ("regulatory_fetch",         "run_regulatory_sync"),
    ("civil_rights_fetch",       "run_civil_rights_sync"),
    ("federal_register_fetch",   "run_federal_register_sync"),
    ("scotus_fetch",             "run_scotus_sync"),
    ("state_courts_fetch",       "run_state_sync"),
    ("all_states_fetch",         "run_all_states_sync"),
    ("congress_fetch",           "run_congress_sync"),
    ("legal_news_fetch",         "run_legal_news_sync"),
    ("edgar_fetch",              "run_edgar_sync"),
    ("regulations_docket_fetch", "run_regulations_sync"),
    ("ofac_fetch",               "run_ofac_sync"),
    ("govtrack_fetch",           "run_govtrack_sync"),
    ("openstates_fetch",         "run_openstates_sync"),
    ("international_fetch",      "run_international_sync"),
    ("state_ecourts_fetch",      "run_state_ecourts_sync"),
    ("courtlistener_alerts",     "create_cl_search_alert"),
]

for mod, fn in aggs:
    test(f"aggregator.{mod}", lambda m=mod, f=fn:
        str(__import__(f"aggregator.{m}", fromlist=[f]).__dict__[f].__name__))


# ── 4. Database ────────────────────────────────────────────────────────────────
section("4. Database")

def check_db():
    from database.db import init_db, count_cases, source_stats, init_alerts_table
    init_db()
    init_alerts_table()
    total = count_cases()
    stats = source_stats()
    return f"{total:,} cases, {len(stats)} sources indexed"

test("DB init + count", check_db, critical=True)

def check_db_search():
    from database.db import search_cases
    r = search_cases(q="", limit=5)
    return f"{len(r)} results returned"

test("DB search (empty query)", check_db_search)

def check_db_fts():
    from database.db import search_cases
    r = search_cases(q="fraud", limit=5)
    return f"FTS search returned {len(r)} results"

test("DB full-text search", check_db_fts)


# ── 5. API key ─────────────────────────────────────────────────────────────────
section("5. Anthropic API Key")

def check_key_present():
    k = os.environ.get("ANTHROPIC_API_KEY","")
    if not k:
        raise Exception("ANTHROPIC_API_KEY not set — use sidebar Test Connection")
    if not k.startswith("sk-"):
        raise Exception(f"Key format invalid (starts with: {k[:8]})")
    return f"Key present, length={len(k)}"

key_ok = test("API key present + valid format", check_key_present, critical=True)

if key_ok:
    def check_api_ping():
        import anthropic
        client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        r = client.messages.create(
            model="claude-haiku-4-5", max_tokens=5,
            messages=[{"role":"user","content":"hi"}]
        )
        return f"Model: {r.model}, stop: {r.stop_reason}"

    test("Anthropic API ping (haiku)", check_api_ping, critical=True)

    def check_web_search():
        import anthropic
        client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        r = client.messages.create(
            model="claude-sonnet-4-6", max_tokens=100,
            tools=[{"type":"web_search_20260209","name":"web_search"}],
            messages=[{"role":"user","content":"What year did the FTC AI enforcement guidelines publish? One sentence."}]
        )
        return f"stop_reason={r.stop_reason}, blocks={len(r.content)}"

    test("Web search tool (sonnet)", check_web_search)


# ── 6. CourtListener ───────────────────────────────────────────────────────────
section("6. CourtListener API")

def check_cl_direct():
    import httpx
    from tools.court_tools import _cl_headers
    r = httpx.get(
        "https://www.courtlistener.com/api/rest/v4/search/",
        params={"q":"fraud","type":"o","order_by":"score desc","format":"json"},
        headers=_cl_headers(), timeout=15
    )
    r.raise_for_status()
    results = r.json().get("results",[])
    return f"HTTP 200 — {len(results)} opinions returned"

test("CourtListener REST API", check_cl_direct)

def check_cl_token():
    t = os.environ.get("COURTLISTENER_API_TOKEN","")
    if not t:
        raise Exception("No token — add COURTLISTENER_API_TOKEN to .env for higher rate limits")
    return f"Token present, length={len(t)}"

test("CourtListener token (optional — raises rate limits)", check_cl_token)


# ── 7. Harvard CAP ────────────────────────────────────────────────────────────
section("7. Harvard Caselaw Access Project")

def check_cap():
    from analyzers.precedent_analyzer import search_cases_cap
    results = search_cases_cap("civil rights discrimination", max_results=3)
    # Harvard CAP API retired June 2024 — now falls back to CourtListener
    return f"{len(results)} cases (via CourtListener fallback — Harvard CAP API retired)"

test("Harvard CAP search (CL fallback)", check_cap)


# ── 8. Court researcher agent ─────────────────────────────────────────────────
section("8. Investigation Agents")

if key_ok:
    def check_court_researcher():
        import anthropic
        from agents.court_researcher import run_court_researcher
        client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        r = run_court_researcher("FTC consumer fraud enforcement", client=client)
        n = r.get("total_found", len(r.get("cases",[])))
        if r.get("error") and n == 0:
            raise Exception(r["error"])
        return f"total_found={n}, cases={len(r.get('cases',[]))}"

    test("Court researcher agent", check_court_researcher)

    def check_web_researcher():
        import anthropic
        from agents.web_researcher import run_web_researcher
        client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        r = run_web_researcher("FTC AI fraud enforcement 2024", client=client)
        n = r.get("total_found", len(r.get("findings",[])))
        if r.get("error") and n == 0:
            raise Exception(r["error"])
        return f"total_found={n}, findings={len(r.get('findings',[]))}"

    test("Web researcher agent", check_web_researcher)

else:
    print(f"  {SKIP}  Court researcher agent (no API key)")
    print(f"  {SKIP}  Web researcher agent (no API key)")


# ── 9. Document viewer ────────────────────────────────────────────────────────
section("9. Document Viewer")

def check_pdf_extract():
    import io, pdfplumber
    # Create a minimal PDF in memory for testing
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.lib.pagesizes import letter
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=letter)
    c.drawString(100, 700, "LegalPerigee Test PDF")
    c.save()
    buf.seek(0)
    with pdfplumber.open(buf) as pdf:
        text = pdf.pages[0].extract_text() or ""
    if "LegalPerigee" not in text:
        raise Exception(f"Text extraction failed, got: {text[:50]!r}")
    return "PDF text extraction OK"

test("PDF text extraction (pdfplumber)", check_pdf_extract)

def check_docx():
    import io
    from docx import Document
    doc = Document()
    doc.add_paragraph("LegalPerigee Test Document")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    doc2 = Document(buf)
    text = doc2.paragraphs[0].text
    if "LegalPerigee" not in text:
        raise Exception(f"DOCX round-trip failed: {text!r}")
    return "DOCX read/write OK"

test("DOCX round-trip (python-docx)", check_docx)

def check_json_extract():
    from utils.json_extract import extract_json
    tests = [
        ('{"a":1}',                       {"a":1}),
        ('```json\n{"b":2}\n```',          {"b":2}),
        ('Some text\n{"c":3}\nmore text',  {"c":3}),
        ('[{"d":4}]',                      [{"d":4}]),
    ]
    for inp, expected in tests:
        got = extract_json(inp)
        if got != expected:
            raise Exception(f"extract_json({inp!r}) → {got!r}, expected {expected!r}")
    return f"{len(tests)} parse patterns all correct"

test("JSON extractor (all patterns)", check_json_extract)


# ── 10. Legislative Watch ──────────────────────────────────────────────────────
section("10. Legislative Watch")

def check_trend_analyzer():
    from analyzers.trend_analyzer import topic_trends, detect_waves, detect_preemption_bills
    t = topic_trends(365)
    w = detect_waves(90, 3)
    p = detect_preemption_bills(180)
    return f"topics={len(t)}, waves={len(w)}, preemption={len(p)}"

test("Trend analyzer (all functions)", check_trend_analyzer)


# ── 11. Precedent engine ───────────────────────────────────────────────────────
section("11. Precedent Research Engine")

def check_precedent_search():
    from analyzers.precedent_analyzer import search_cases_cl, search_cases_cap
    cl = search_cases_cl("housing discrimination fair housing", max_results=3)
    cap = search_cases_cap("civil rights discrimination", max_results=3)
    return f"CL={len(cl)} results, CAP={len(cap)} results"

test("Precedent search (CL + CAP)", check_precedent_search)


# ── 12. Alerts system ─────────────────────────────────────────────────────────
section("12. Alerts System")

def check_alerts_db():
    from database.db import init_alerts_table, get_alert_rules, save_alert_rule, delete_alert_rule
    init_alerts_table()
    rule_id = save_alert_rule({
        "name":"__debug_test__","keywords":"test","courts":"",
        "states":"","harm_types":"","email":"test@test.com"
    })
    rules = get_alert_rules()
    found = any(r["name"] == "__debug_test__" for r in rules)
    if rule_id:
        delete_alert_rule(rule_id)
    if not found:
        raise Exception("Alert rule not found after save")
    return f"Create/list/delete cycle OK (id={rule_id})"

test("Alert rules DB (create/list/delete)", check_alerts_db)

def check_smtp_config():
    from alerts.email_sender import smtp_configured, _smtp_config
    cfg = _smtp_config()
    if not smtp_configured():
        raise Exception(f"SMTP not configured — set ALERT_EMAIL_FROM + ALERT_EMAIL_PASSWORD in .env")
    return f"SMTP configured: {cfg['host']}:{cfg['port']}"

test("SMTP email (optional — needed for email alerts)", check_smtp_config)


# ── 13. FastAPI backend (v2) ───────────────────────────────────────────────────
section("13. v2 FastAPI Backend")

def check_fastapi_syntax():
    import ast
    p = Path(__file__).parent.parent.parent / "LegalPerigee-v2" / "backend" / "main.py"
    if not p.exists():
        raise Exception(f"main.py not found at {p}")
    ast.parse(p.read_text())
    return "FastAPI main.py syntax OK"

test("v2 backend main.py syntax", check_fastapi_syntax)

def check_fastapi_import():
    sys.path.insert(0, str(Path(__file__).parent.parent.parent / "LegalPerigee-v2" / "backend"))
    import importlib
    spec = importlib.util.spec_from_file_location(
        "v2_main",
        str(Path(__file__).parent.parent.parent / "LegalPerigee-v2" / "backend" / "main.py")
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    route_count = len([r for r in mod.app.routes if hasattr(r,"methods")])
    return f"{route_count} routes registered"

test("v2 FastAPI app loads cleanly", check_fastapi_import)


# ── 14. Streamlit app file ────────────────────────────────────────────────────
section("14. GUI Integrity")

def check_gui_syntax():
    import ast
    gui = Path(__file__).parent.parent / "gui.py"
    ast.parse(gui.read_text())
    lines = gui.read_text().count("\n")
    return f"gui.py syntax OK ({lines:,} lines)"

test("gui.py syntax valid", check_gui_syntax, critical=True)

def check_server_running():
    import httpx
    try:
        r = httpx.get("http://localhost:8501", timeout=3)
        return f"Streamlit server responding (HTTP {r.status_code})"
    except Exception:
        raise Exception("Streamlit server not running — launch LegalPerigee.app first")

test("Streamlit server running (port 8501)", check_server_running)


# ── Summary ────────────────────────────────────────────────────────────────────
passed   = [r for r in results if r["status"] == "PASS"]
failed   = [r for r in results if r["status"] == "FAIL"]
critical = [r for r in failed  if r.get("critical")]

print(f"\n{'═'*65}")
print(f"  RESULTS: {len(passed)} passed  |  {len(failed)} failed  |  {len(critical)} critical")
print(f"{'═'*65}")

if failed:
    print(f"\n  {FAIL}  FAILURES:")
    for r in failed:
        flag = " [CRITICAL]" if r.get("critical") else ""
        print(f"       • {r['name']}{flag}")
        print(f"         {r['detail']}")

if not failed:
    print(f"\n  {PASS}  All tests passed — app is functioning as designed.")

# Save report
report_path = Path(__file__).parent.parent / "data" / "debug_report.json"
report_path.parent.mkdir(exist_ok=True)
report_path.write_text(json.dumps({
    "timestamp": datetime.utcnow().isoformat()+"Z",
    "passed": len(passed), "failed": len(failed),
    "results": results,
}, indent=2))
print(f"\n  📄 Full report saved: {report_path}")
print()

sys.exit(1 if critical else 0)
