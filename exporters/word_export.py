"""
Export a CaseIntelReport as a formatted Word (.docx) legal memo.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from models.case_report import CaseIntelReport

NAVY   = RGBColor(0x1a, 0x27, 0x44)
GOLD   = RGBColor(0xC9, 0xA8, 0x4C)
RED    = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xD6, 0x7F, 0x1E)
GREEN  = RGBColor(0x27, 0xAE, 0x60)
GRAY   = RGBColor(0x55, 0x55, 0x55)

SEV_COLORS = {"high": RED, "medium": ORANGE, "low": GREEN, "unknown": GRAY}


def _heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY


def _para(doc: Document, text: str, bold: bool = False, color: RGBColor = None,
          size_pt: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color


def _kv(doc: Document, key: str, value: str) -> None:
    p = doc.add_paragraph()
    r1 = p.add_run(f"{key}: ")
    r1.bold = True
    r1.font.color.rgb = NAVY
    r1.font.size = Pt(10)
    r2 = p.add_run(value or "—")
    r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(2)


def export_to_docx(report: CaseIntelReport) -> bytes:
    """Generate a Word document from a CaseIntelReport. Returns bytes."""
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # ── Cover header ──────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("⚖  LEGALPERIGEE")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("CASE INTELLIGENCE REPORT")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = GOLD
    sub_run.bold = True

    doc.add_paragraph()  # spacer

    # ── Meta block ────────────────────────────────────────────────────────────
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    cells = [
        ("Query",     report.query),
        ("Generated", report.generated_at[:19].replace("T", " ")),
        ("Total Cases", str(report.total_cases_found)),
        ("High Severity", str(report.high_severity_count)),
    ]
    for i, (k, v) in enumerate(cells):
        meta_table.rows[i].cells[0].paragraphs[0].add_run(k).bold = True
        meta_table.rows[i].cells[1].paragraphs[0].add_run(v)

    doc.add_paragraph()

    # ── Executive Summary ─────────────────────────────────────────────────────
    _heading(doc, "Executive Summary", level=1)
    doc.add_paragraph(report.summary)

    if report.filters_applied and not report.filters_applied.is_empty():
        _heading(doc, "Search Filters Applied", level=2)
        doc.add_paragraph(report.filters_applied.as_query_string())

    doc.add_page_break()

    # ── Case Findings ─────────────────────────────────────────────────────────
    _heading(doc, "Case Findings", level=1)

    for i, finding in enumerate(report.findings, 1):
        sev = finding.severity.value
        color = SEV_COLORS.get(sev, GRAY)

        # Case header
        case_p = doc.add_paragraph()
        num_run = case_p.add_run(f"{i}. ")
        num_run.bold = True
        num_run.font.color.rgb = NAVY
        name_run = case_p.add_run(finding.case_name)
        name_run.bold = True
        name_run.font.size = Pt(13)
        name_run.font.color.rgb = NAVY
        sev_run = case_p.add_run(f"  [{sev.upper()}]")
        sev_run.bold = True
        sev_run.font.color.rgb = color

        # Metadata
        _kv(doc, "Court",            finding.court_name or finding.jurisdiction or "—")
        _kv(doc, "Filing Date",      finding.filing_date or "Unknown")
        _kv(doc, "Legal Status",     finding.legal_status or "Unknown")
        _kv(doc, "AI Actor",         finding.ai_actor or "Unknown")
        _kv(doc, "Harm Types",       ", ".join(h.value for h in finding.harm_types) or "—")
        if finding.protected_classes:
            _kv(doc, "Protected Classes",
                ", ".join(p.value for p in finding.protected_classes))

        # Main text
        doc.add_paragraph()
        _para(doc, "Victims:", bold=True)
        doc.add_paragraph(finding.victim_description)
        _para(doc, "Practice Alleged:", bold=True)
        doc.add_paragraph(finding.deceptive_practice)
        _para(doc, "Verifiable Harm:", bold=True)
        doc.add_paragraph(finding.verifiable_harm)

        # Sources
        if finding.sources:
            _para(doc, "Sources:", bold=True)
            for src in finding.sources:
                link = src.url or "—"
                doc.add_paragraph(
                    f"• [{src.source_type}] {src.title} ({src.date or ''})\n  {link}",
                    style="List Bullet",
                )

        doc.add_paragraph()
        doc.add_paragraph("─" * 60)
        doc.add_paragraph()

    # ── Investigator Notes ────────────────────────────────────────────────────
    doc.add_page_break()
    _heading(doc, "Investigator Notes", level=1)
    doc.add_paragraph(report.investigator_notes)

    # ── Sources Searched ──────────────────────────────────────────────────────
    if report.sources_searched:
        _heading(doc, "Sources Searched", level=2)
        for s in report.sources_searched:
            doc.add_paragraph(f"• {s}", style="List Bullet")

    # ── Footer ────────────────────────────────────────────────────────────────
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run(
        f"CONFIDENTIAL — LegalPerigee · Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    fr.font.size = Pt(8)
    fr.font.color.rgb = GRAY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
